"""Workbook login and server-side session/role enforcement."""
import os, time, secrets, hashlib, threading
from collections import defaultdict
from pathlib import Path
from urllib.parse import parse_qs
from openpyxl import load_workbook
from fastapi import Request
from fastapi.responses import JSONResponse, RedirectResponse, FileResponse
from starlette.middleware.base import BaseHTTPMiddleware

LIMITED_PAGES = ['enquiry', 'followup', 'testdrive', 'booking', 'sales', 'exchange']
FULL_PAGES = ['overview', *LIMITED_PAGES, 'inventory', 'stock', 'comparison']
DATA_FILE = Path(os.environ.get('AUTH_DATA_FILE', str(Path(__file__).parent / 'data' / 'Data.xlsx')))
SESSIONS = {}
FAILURES = defaultdict(list)
LOCK = threading.Lock()
TTL = 8 * 60 * 60

def accounts():
    """Read afresh so account removal, password and role changes revoke sessions."""
    if not DATA_FILE.exists():
        raise ValueError('Place your user account workbook at data/Data.xlsx.')
    wb = load_workbook(DATA_FILE, read_only=True, data_only=True)
    try:
        ws = wb['Users'] if 'Users' in wb.sheetnames else wb.active
        rows = iter(ws.values)
        headers = [str(v or '').strip().lower().replace('_', ' ') for v in next(rows)]
        aliases = {'userid':'user id', 'username':'user id', 'employee id':'user id', 'access level':'access', 'full/limited':'access'}
        headers = [aliases.get(h, h) for h in headers]
        if not {'user id','password','access'}.issubset(headers):
            raise ValueError('Users sheet needs User ID, Password and Access columns.')
        result = {}
        for values in rows:
            row = dict(zip(headers, [str(v).strip() if v is not None else '' for v in values]))
            uid = row.get('user id', '')
            if not uid: continue
            role = row['access'].lower().replace(' access', '').strip()
            if role not in ('full', 'limited') or not row['password'] or uid in result:
                raise ValueError('Invalid or duplicate account in Users sheet.')
            if row.get('active','yes').lower() in ('no','false','0','inactive'): continue
            row['access'] = role
            row['revision'] = hashlib.sha256(repr(sorted(row.items())).encode()).hexdigest()
            result[uid] = row
        return result
    finally:
        wb.close()

def identity(request):
    token = request.cookies.get('unnati_session', '')
    with LOCK:
        session = SESSIONS.get(token)
    if not session or session['expires'] <= time.time(): return None
    try: user = accounts().get(session['uid'])
    except (ValueError, OSError, StopIteration): return None
    if not user or user['revision'] != session['revision']: return None
    return user

def public_user(user):
    return {'user_id':user['user id'], 'name':user.get('name') or user['user id'], 'access':user['access'],
            'pages':FULL_PAGES if user['access']=='full' else LIMITED_PAGES}

def install_auth(app, static_dir):
    @app.get('/login')
    def login_page():
        return FileResponse(Path(static_dir)/'login.html', headers={'Cache-Control':'no-store'})

    @app.post('/api/auth/login')
    async def login(request: Request):
        if request.headers.get('origin') and request.headers['origin'] != str(request.base_url).rstrip('/'):
            return JSONResponse({'detail':'Invalid origin'}, status_code=403)
        body = await request.body()
        if len(body) > 8192: return JSONResponse({'detail':'Request too large'}, status_code=413)
        form = parse_qs(body.decode('utf-8', errors='replace'))
        uid = form.get('user_id',[''])[0].strip()
        password = form.get('password',[''])[0]
        key = (request.client.host if request.client else '', uid)
        now = time.time()
        with LOCK:
            FAILURES[key] = [t for t in FAILURES[key] if now-t < 900]
            if len(FAILURES[key]) >= 10:
                return JSONResponse({'detail':'Too many attempts. Try again in 15 minutes.'}, status_code=429)
        try: user = accounts().get(uid)
        except (ValueError, OSError, StopIteration) as exc:
            return JSONResponse({'detail':str(exc)}, status_code=503)
        if not user or not secrets.compare_digest(password.encode(),user['password'].encode()):
            with LOCK: FAILURES[key].append(now)
            return JSONResponse({'detail':'User ID or password is incorrect.'}, status_code=401)
        token = secrets.token_urlsafe(32)
        with LOCK:
            for old in list(SESSIONS):
                if SESSIONS[old]['expires'] <= now: SESSIONS.pop(old,None)
            SESSIONS.pop(request.cookies.get('unnati_session',''),None)
            SESSIONS[token] = {'uid':uid,'revision':user['revision'],'expires':now+TTL}
            FAILURES.pop(key,None)
        response = JSONResponse(public_user(user))
        response.set_cookie('unnati_session',token,max_age=TTL,httponly=True,samesite='strict',
                            secure=os.environ.get('COOKIE_SECURE','0')=='1' or request.url.scheme=='https')
        return response

    @app.get('/api/auth/me')
    def me(request: Request):
        return public_user(request.state.user)

    @app.get('/logout')
    def logout(request: Request):
        with LOCK: SESSIONS.pop(request.cookies.get('unnati_session',''),None)
        response = RedirectResponse('/login',status_code=303)
        response.delete_cookie('unnati_session')
        return response

    app.add_middleware(AccessMiddleware)

class AccessMiddleware(BaseHTTPMiddleware):
    async def dispatch(self, request, call_next):
        path = request.url.path
        public = path in ('/login','/api/auth/login','/api/health','/logo.png')
        if public: return await call_next(request)
        user = identity(request)
        if not user:
            if path.startswith('/api/'):
                return JSONResponse({'detail':'Please sign in.'}, status_code=401)
            return RedirectResponse('/login',status_code=303)
        request.state.user = user
        if request.method not in ('GET','HEAD','OPTIONS'):
            origin = request.headers.get('origin')
            if origin and origin != str(request.base_url).rstrip('/'):
                return JSONResponse({'detail':'Invalid origin'},status_code=403)
        if user['access']=='limited' and path.startswith('/api/'):
            allowed = {'/api/auth/me','/api/meta','/api/filters','/api/kpis','/api/enquiry','/api/test-drive',
                       '/api/booking','/api/sales','/api/exchange','/api/followup','/api/followup/list',
                       '/api/followup/booked','/api/breakdown'}
            if request.method != 'GET' or path not in allowed or (path=='/api/breakdown' and request.query_params.get('section') not in LIMITED_PAGES):
                return JSONResponse({'detail':'This screen is not available for Limited access.'}, status_code=403)
        response = await call_next(request)
        response.headers['Cache-Control'] = 'no-store'
        return response
